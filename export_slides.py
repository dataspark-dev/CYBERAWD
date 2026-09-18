"""
Export the Synergy Cyber Security Awareness Month deck to a .pptx file.

Renders each slides/slide-NN.html file with a headless browser at
1920x1080 and drops the screenshot into a 16:9 PowerPoint deck, one
slide per screenshot. Read-only against the deck itself — it only
reads scripts/deck.js to find out how many slides exist and in what
order, the same source of truth extract_deck_manifest.py uses, so it
auto-adapts whenever slides are added, removed, or reordered.

Usage:
    python app.py                    # in one terminal, start the server
    python export_slides.py          # in another terminal, run the export

Requires: playwright (with `playwright install chromium` run once),
python-pptx. Both already available in this project's environment.
"""

import asyncio
import os
import re
import shutil
import sys
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches


# ── Configuration ────────────────────────────────────────────────────────────

ROOT            = Path(__file__).resolve().parent
DECK_JS         = ROOT / "scripts" / "deck.js"

# The server this hits — start it first with `python app.py` (defaults to
# port 5000; override with the PORT env var the same way app.py does), or
# point this at a plain `python -m http.server <port>` run from the repo
# root, since the slides are static files either way.
SERVER_BASE     = os.environ.get("EXPORT_BASE_URL", "http://localhost:5000")
OUTPUT_FILE     = "Synergy_Cyber_Security_Awareness_Month.pptx"

# Canvas size in CSS pixels — matches body/html in every slides/slide-NN.html
VIEWPORT_WIDTH  = 1920
VIEWPORT_HEIGHT = 1080

# ── SCALE FACTOR ─────────────────────────────────────────────────────────────
# 2 → 3840 x 2160 px per slide (~288 DPI in PPTX). Crisp on projectors/screens.
# 3 → 5760 x 3240 px per slide (~432 DPI). Use if you need ultra-sharp print.
# NOTE: scale=3 increases RAM usage and export time significantly.
SCALE           = 2

# ── Derived constants (do not edit) ──────────────────────────────────────────
# True 16:9 widescreen standard for PowerPoint / Google Slides
PPTX_WIDTH_IN   = Inches(13.333)
PPTX_HEIGHT_IN  = Inches(7.5)

TEMP_DIR        = Path("./temp_slides")
KEEP_TEMP       = "--keep-temp" in sys.argv


def get_slide_files():
    """Read scripts/deck.js's SLIDES array — same parsing approach app.py
    uses at startup to validate slides/ against the deck.js manifest, and
    extract_deck_manifest.py uses to build deck-manifest.json/md. Keeps
    this script from ever hardcoding a slide count that drifts out of date."""
    text = DECK_JS.read_text(encoding="utf-8")
    match = re.search(r"(?:var|const)\s+SLIDES\s*=\s*\[(.*?)\]\s*;", text, re.DOTALL)
    if not match:
        print("ERROR: could not find SLIDES array in scripts/deck.js", file=sys.stderr)
        sys.exit(1)
    files = re.findall(r"file:\s*'([^']+)'", match.group(1))
    if not files:
        print("ERROR: SLIDES array parsed but no slide files found", file=sys.stderr)
        sys.exit(1)
    missing = [f for f in files if not (ROOT / "slides" / f).exists()]
    if missing:
        print(f"ERROR: deck.js references file(s) missing from slides/: {missing}", file=sys.stderr)
        sys.exit(1)
    return files


# ── Playwright init script injected into every page ──────────────────────────
INIT_SCRIPT = """
window.addEventListener('DOMContentLoaded', () => {
    const style = document.createElement('style');
    style.textContent = `
        *, *::before, *::after {
            animation-duration: 0s !important;
            transition-duration: 0s !important;
            animation-delay:    0s !important;
        }
        html, body {
            overflow: hidden !important;
            margin:   0 !important;
            padding:  0 !important;
            -webkit-font-smoothing: antialiased !important;
            text-rendering: optimizeLegibility !important;
        }
        ::-webkit-scrollbar { display: none !important; }
    `;
    document.head.appendChild(style);
});
"""


async def export_to_pptx():
    slide_files = get_slide_files()
    total = len(slide_files)
    print(f"deck.js declares {total} slide(s): {slide_files[0]} .. {slide_files[-1]}")

    TEMP_DIR.mkdir(exist_ok=True)

    prs = Presentation()
    prs.slide_width  = PPTX_WIDTH_IN
    prs.slide_height = PPTX_HEIGHT_IN

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            args=[
                "--font-render-hinting=none",
                "--disable-lcd-text",           # greyscale AA -> sharper at HiDPI
                "--disable-dev-shm-usage",      # avoids /dev/shm OOM on Linux
                "--no-sandbox",
            ]
        )

        context = await browser.new_context(
            viewport={"width": VIEWPORT_WIDTH, "height": VIEWPORT_HEIGHT},
            device_scale_factor=SCALE,          # single source of truth for DPI
        )
        await context.add_init_script(INIT_SCRIPT)

        page = await context.new_page()

        for idx, filename in enumerate(slide_files, start=1):
            slide_url = f"{SERVER_BASE}/slides/{filename}"
            pct = (idx - 1) / total * 100
            print(f"  [{pct:5.1f}%]  Capturing slide {idx:02d}/{total}  ...  {slide_url}")

            try:
                await page.goto(slide_url, wait_until="networkidle", timeout=30_000)

                # Wait for web fonts to finish loading
                await page.evaluate("() => document.fonts.ready")

                # Force all <img> elements to finish loading
                await page.evaluate("""() => Promise.all(
                    Array.from(document.images)
                         .filter(img => !img.complete)
                         .map(img => new Promise(resolve => {
                             img.onload  = resolve;
                             img.onerror = resolve;
                         }))
                )""")

                # Extra settle time for CSS entrance animations (.reveal-stagger)
                await asyncio.sleep(1.0)

                image_path = TEMP_DIR / f"slide_{idx:02d}.png"

                # clip uses CSS pixels - Playwright handles the device_scale_factor
                # internally and writes a (width*SCALE) x (height*SCALE) file
                # automatically. Do NOT multiply clip by SCALE here.
                await page.screenshot(
                    path=str(image_path),
                    clip={
                        "x":      0,
                        "y":      0,
                        "width":  VIEWPORT_WIDTH,
                        "height": VIEWPORT_HEIGHT,
                    },
                    full_page=False,
                    animations="disabled",
                    type="png",                     # lossless - no JPEG artefacts
                )

                # ── Add to PowerPoint ────────────────────────────────────────
                slide_layout = prs.slide_layouts[6]   # blank layout
                slide = prs.slides.add_slide(slide_layout)
                slide.shapes.add_picture(
                    str(image_path),
                    Inches(0), Inches(0),
                    width=prs.slide_width,
                    height=prs.slide_height,
                )

                print(f"           done  Slide {idx:02d} captured "
                      f"({VIEWPORT_WIDTH * SCALE}x{VIEWPORT_HEIGHT * SCALE} px)")

            except Exception as e:
                print(f"           FAILED on {filename}: {e}")
                # Save a debug screenshot at native resolution for inspection
                try:
                    err_path = TEMP_DIR / f"error_{filename.replace('.html', '')}.png"
                    await page.screenshot(path=str(err_path), full_page=False)
                    print(f"             Debug screenshot saved -> {err_path}")
                except Exception:
                    pass

        await browser.close()

    prs.save(OUTPUT_FILE)

    # ── Cleanup ──────────────────────────────────────────────────────────────
    if not KEEP_TEMP:
        shutil.rmtree(TEMP_DIR, ignore_errors=True)
        print("  Temp files cleaned up.")
    else:
        print(f"  Temp files retained in: {TEMP_DIR.resolve()}")

    print("-" * 60)
    print(f"Saved:  {OUTPUT_FILE}")
    print(f"   Server:           {SERVER_BASE}")
    print(f"   Slide canvas:     {VIEWPORT_WIDTH} x {VIEWPORT_HEIGHT} CSS px")
    print(f"   Scale factor:     {SCALE}x")
    print(f"   Screenshot size:  {VIEWPORT_WIDTH * SCALE} x {VIEWPORT_HEIGHT * SCALE} px")
    print(f"   PPTX slide size:  13.333\" x 7.5\"  (true 16:9 widescreen)")
    eff_dpi = round(VIEWPORT_WIDTH * SCALE / 13.333)
    print(f"   Effective DPI:    ~{eff_dpi} DPI")
    print(f"   Total slides:     {total}")


if __name__ == "__main__":
    asyncio.run(export_to_pptx())
