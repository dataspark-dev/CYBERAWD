"""
Export the de-branded, industry-neutral general deck (slides-general-clean/)
to a .pptx file.

This is the "no logo, no Synergy/AFT references, no maritime-specific
content" variant of the 19-slide general deck, built for a prospective new
client whose workforce isn't maritime-specific. It lives in
slides-general-clean/ (a sibling of slides/, not inside it, so its own
../styles, ../assets, ../scripts relative paths keep resolving to the same
shared repo-root folders) and is read-only against the live Synergy-branded
deck - nothing under slides/ is touched by this script or by the one that
generated slides-general-clean/.

Renders each slides-general-clean/slide-NN.html file directly via a file://
URL (no Flask server needed - these are static pages with no fetch() calls)
with a headless browser at 1920x1080 and drops the screenshot into a 16:9
PowerPoint deck, one slide per screenshot. Same capture approach as
export_slides.py/export_audience_slides.py.

Usage:
    python export_general_clean_slides.py

Requires: playwright (with `playwright install chromium` run once),
python-pptx. Both already available in this project's environment.
"""

import asyncio
import shutil
import sys
from pathlib import Path
from urllib.request import pathname2url

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

from export_slides import (
    INIT_SCRIPT,
    PPTX_HEIGHT_IN,
    PPTX_WIDTH_IN,
    SCALE,
    VIEWPORT_HEIGHT,
    VIEWPORT_WIDTH,
)

ROOT = Path(__file__).resolve().parent
SLIDES_DIR = ROOT / "slides-general-clean"
OUTPUT_FILE = "Cyber_Security_Awareness_Program_General_NoLogo.pptx"
TEMP_DIR = ROOT / "temp_slides_general_clean"


def get_slide_files():
    files = sorted(SLIDES_DIR.glob("slide-*.html"))
    if not files:
        print(f"ERROR: no slide-*.html files found in {SLIDES_DIR}", file=sys.stderr)
        sys.exit(1)
    return [f.name for f in files]


async def export_to_pptx():
    slide_files = get_slide_files()
    total = len(slide_files)
    print(f"{SLIDES_DIR.name} contains {total} slide(s): {slide_files[0]} .. {slide_files[-1]}")

    TEMP_DIR.mkdir(exist_ok=True)

    prs = Presentation()
    prs.slide_width = PPTX_WIDTH_IN
    prs.slide_height = PPTX_HEIGHT_IN

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            args=[
                "--font-render-hinting=none",
                "--disable-lcd-text",
                "--disable-dev-shm-usage",
                "--no-sandbox",
            ]
        )

        context = await browser.new_context(
            viewport={"width": VIEWPORT_WIDTH, "height": VIEWPORT_HEIGHT},
            device_scale_factor=SCALE,
        )
        await context.add_init_script(INIT_SCRIPT)

        page = await context.new_page()

        for idx, filename in enumerate(slide_files, start=1):
            file_path = SLIDES_DIR / filename
            slide_url = "file://" + pathname2url(str(file_path))
            pct = (idx - 1) / total * 100
            print(f"  [{pct:5.1f}%]  Capturing slide {idx:02d}/{total}  ...  {filename}")

            try:
                await page.goto(slide_url, wait_until="networkidle", timeout=30_000)
                await page.evaluate("() => document.fonts.ready")
                await page.evaluate("""() => Promise.all(
                    Array.from(document.images)
                         .filter(img => !img.complete)
                         .map(img => new Promise(resolve => {
                             img.onload  = resolve;
                             img.onerror = resolve;
                         }))
                )""")
                await asyncio.sleep(1.0)

                image_path = TEMP_DIR / f"slide_{idx:02d}.png"
                await page.screenshot(
                    path=str(image_path),
                    clip={"x": 0, "y": 0, "width": VIEWPORT_WIDTH, "height": VIEWPORT_HEIGHT},
                    full_page=False,
                    animations="disabled",
                    type="png",
                )

                slide_layout = prs.slide_layouts[6]
                slide = prs.slides.add_slide(slide_layout)
                slide.shapes.add_picture(
                    str(image_path), Inches(0), Inches(0),
                    width=prs.slide_width, height=prs.slide_height,
                )

                print(f"           done  Slide {idx:02d} captured")

            except Exception as e:
                print(f"           FAILED on {filename}: {e}")
                try:
                    err_path = TEMP_DIR / f"error_{filename.replace('.html', '')}.png"
                    await page.screenshot(path=str(err_path), full_page=False)
                    print(f"             Debug screenshot saved -> {err_path}")
                except Exception:
                    pass

        await browser.close()

    prs.save(OUTPUT_FILE)
    shutil.rmtree(TEMP_DIR, ignore_errors=True)
    print(f"Saved: {OUTPUT_FILE}  ({total} slides)")


if __name__ == "__main__":
    asyncio.run(export_to_pptx())
